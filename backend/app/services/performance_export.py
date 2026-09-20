"""Excel export for performance analytics (probability calibration)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True)


def _write_table(ws: Worksheet, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    for index, header in enumerate(headers, start=1):
        width = max(len(str(header)), *(len(str(row[index - 1])) for row in rows)) if rows else len(str(header))
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 40)
    if rows:
        ws.freeze_panes = "A2"


def probability_calibration_workbook(data: dict) -> io.BytesIO:
    """Build an .xlsx workbook mirroring the probability-calibration analytics view."""
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    s = data["summary"]
    _write_table(
        summary_ws,
        ["Metric", "Value"],
        [
            ["Model version", data["model_version"]],
            ["Sample size", s["sample_size"]],
            ["Voids excluded", s["voids_excluded"]],
            ["Unsupported markets excluded", s["unsupported_markets_excluded"]],
            ["Lohela Brier score", s["lohela_brier_score"]],
            ["Market Brier score", s["market_brier_score"]],
            ["Lohela calibration error", s["lohela_calibration_error"]],
            ["Market calibration error", s["market_calibration_error"]],
            ["Note", data["note"]],
        ],
    )

    bucket_headers = ["Band", "Sample size", "Wins", "Losses", "Hit rate", "Avg probability", "Calibration gap", "ROI"]

    lohela_ws = wb.create_sheet("Lohela buckets")
    _write_table(lohela_ws, bucket_headers, [
        [r["band"], r["sample_size"], r["wins"], r["losses"], r["hit_rate"], r["avg_probability"], r["calibration_gap"], r["roi"]]
        for r in data["lohela_buckets"]
    ])

    market_ws = wb.create_sheet("Market buckets")
    _write_table(market_ws, bucket_headers, [
        [r["band"], r["sample_size"], r["wins"], r["losses"], r["hit_rate"], r["avg_probability"], r["calibration_gap"], r["roi"]]
        for r in data["market_buckets"]
    ])

    matrix_headers = ["Lohela band", "Market band", "Sample size", "Wins", "Losses", "Hit rate", "ROI", "Avg Lohela probability", "Avg market probability"]

    matrix_ws = wb.create_sheet("Combined matrix")
    _write_table(matrix_ws, matrix_headers, [
        [r["lohela_band"], r["market_band"], r["sample_size"], r["wins"], r["losses"], r["hit_rate"], r["roi"], r["avg_lohela_probability"], r["avg_market_probability"]]
        for r in data["combined_matrix"]
    ])

    best_ws = wb.create_sheet("Best combinations")
    _write_table(best_ws, matrix_headers, [
        [r["lohela_band"], r["market_band"], r["sample_size"], r["wins"], r["losses"], r["hit_rate"], r["roi"], r["avg_lohela_probability"], r["avg_market_probability"]]
        for r in data["best_combinations"]
    ])

    edge_headers = ["Edge band", "Sample size", "Wins", "Losses", "Hit rate", "ROI"]
    edge_ws = wb.create_sheet("Edge buckets")
    _write_table(edge_ws, edge_headers, [
        [r["edge_band"], r["sample_size"], r["wins"], r["losses"], r["hit_rate"], r["roi"]]
        for r in data["edge_buckets"]
    ])

    market_headers = [
        "Market family", "Sample size", "Wins", "Losses", "Hit rate", "ROI",
        "Lohela Brier", "Market Brier", "Lohela calibration error", "Market calibration error",
    ]
    market_family_ws = wb.create_sheet("By market")
    _write_table(market_family_ws, market_headers, [
        [
            r["market_family"], r["sample_size"], r["wins"], r["losses"], r["hit_rate"], r["roi"],
            r["lohela_brier_score"], r["market_brier_score"], r["lohela_calibration_error"], r["market_calibration_error"],
        ]
        for r in data["by_market"]
    ])

    edge_by_market_ws = wb.create_sheet("By market edge bands")
    edge_by_market_rows = []
    for market_row in data["by_market"]:
        for edge_row in market_row["edge_buckets"]:
            edge_by_market_rows.append([
                market_row["market_family"], edge_row["edge_band"], edge_row["sample_size"],
                edge_row["wins"], edge_row["losses"], edge_row["hit_rate"], edge_row["roi"],
            ])
    _write_table(edge_by_market_ws, ["Market family", "Edge band", "Sample size", "Wins", "Losses", "Hit rate", "ROI"], edge_by_market_rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
